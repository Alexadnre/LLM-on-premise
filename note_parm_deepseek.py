from openpyxl import load_workbook
from langchain_ollama import ChatOllama
from tqdm import tqdm
import os

# Config LLM
llm = ChatOllama(
    model="deepseek-llm",
    device="gpu",
    trust_remote_code=True,
    temperature=0.2,
    top_k=5,
    top_p=0.9,
    max_new_tokens=200
)

EXCEL_PATH = "resultats_tests.xlsx"

def score_response_no_context(question, response):
    prompt = f"""
Tu es un expert chargé d’évaluer une réponse produite par une IA. 
Voici la question :
"{question}"

Voici la réponse générée :
"{response}"

Note cette réponse sur 100, selon :
- Pertinence
- Clarté
- Cohérence

Réponds uniquement avec un nombre entier.
"""
    try:
        score = llm.invoke(prompt).content.strip()
        return int(score)
    except:
        return None

def add_auto_scores_to_excel(filename):
    if not os.path.exists(filename):
        print(f"❌ Fichier introuvable : {filename}")
        return

    wb = load_workbook(filename)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        print(f"\n🔍 Feuille '{sheet_name}' : Headers = {headers}")

        if "Question" not in headers or "Réponse" not in headers:
            print(f"⛔️ Feuille '{sheet_name}' ignorée (pas de colonne 'Question' ou 'Réponse')")
            continue

        question_idx = headers.index("Question") + 1
        response_idx = headers.index("Réponse") + 1

        # Ajouter la colonne de note si besoin
        if "Note (/100)" not in headers:
            note_col_idx = len(headers) + 1
            ws.cell(row=1, column=note_col_idx, value="Note (/100)")
        else:
            note_col_idx = headers.index("Note (/100)") + 1

        print(f"📝 Évaluation en cours sur la feuille : {sheet_name}")
        for row in tqdm(range(2, ws.max_row + 1), desc=sheet_name):
            if ws.cell(row=row, column=note_col_idx).value is not None:
                continue  # Ne pas écraser les notes déjà présentes

            question = ws.cell(row=row, column=question_idx).value
            response = ws.cell(row=row, column=response_idx).value
            if not question or not response:
                continue

            score = score_response_no_context(question, response)
            if score is not None:
                ws.cell(row=row, column=note_col_idx, value=score)

    wb.save(filename)
    print("\n✅ Toutes les notes ont été ajoutées dans le fichier Excel.")

# Lancer le traitement
add_auto_scores_to_excel(EXCEL_PATH)
